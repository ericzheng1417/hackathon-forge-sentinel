"""Distribution checks (Anderson-Darling, Q-Q, expected ppm) and X̄-S charts for coupon subgroups.
Run: python -m unittest discover tests"""
import io
import json
import math
import pathlib
import subprocess
import sys
import unittest

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from forge_sentinel import data, engine, exports, views  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

# 40 values from N(950, 12); A² = 0.31847 computed independently with scipy.stats.anderson
NORMAL_SAMPLE = [946.9, 956.1, 947.3, 946.2, 938.8, 947.4, 963.3, 955.1, 962.4, 953.0, 954.7, 952.2, 930.0, 960.3,
                 956.1, 956.0, 929.7, 929.1, 939.3, 944.4, 953.7, 949.4, 956.3, 942.3, 953.7, 954.7, 942.1, 970.6,
                 956.7, 964.4, 942.6, 941.1, 945.9, 948.7, 957.6, 953.0, 944.6, 938.5, 943.8, 964.7]
COUPONS = data.parse_csv((ROOT / "datasets" / "04_coupon_spread.csv").read_bytes())


class Normality(unittest.TestCase):
    def test_anderson_darling_matches_scipy(self):
        r = engine.anderson_darling(NORMAL_SAMPLE)
        self.assertAlmostEqual(r["ad"], 0.3184718740509709, places=9)
        self.assertAlmostEqual(r["ad_adj"], r["ad"] * (1 + 0.75 / 40 + 2.25 / 1600))
        self.assertGreater(r["p"], 0.05)
        self.assertTrue(r["normal"])

    def test_detects_non_normal(self):
        skewed = [math.exp(k / 10) for k in range(40)]
        r = engine.anderson_darling(skewed)
        self.assertLess(r["p"], 0.05)
        self.assertFalse(r["normal"])

    def test_small_or_constant_samples(self):
        self.assertIsNone(engine.anderson_darling([1, 2, 3])["normal"])
        self.assertIsNone(engine.anderson_darling([5.0] * 20)["normal"])

    def test_quantiles_and_expected_ppm(self):
        z = engine.normal_quantiles(9)
        self.assertAlmostEqual(z[4], 0.0)
        self.assertAlmostEqual(z[0], -z[-1])
        e = engine.expected_out_of_spec(950, 15, 900, None)
        self.assertAlmostEqual(e["below_ppm"], 429.1, places=1)     # Φ(−3.333) × 10⁶
        e2 = engine.expected_out_of_spec(1000, 30, 910, 1090)
        self.assertAlmostEqual(e2["below_ppm"], e2["above_ppm"])
        self.assertAlmostEqual(e2["total_ppm"], 2699.8, places=1)   # 2 × Φ(−3)

    def test_strength_analysis_carries_distribution(self):
        rows = data.parse_csv((ROOT / "datasets" / "01_stable_capable.csv").read_bytes())
        a = engine.analyze_stream([r for r in rows if r["machine"] == "LPBF-A"], usl=1100)
        self.assertTrue(a.normal["normal"])
        self.assertIn("overall", a.expected)


class XbarSConstants(unittest.TestCase):
    TABLE = {2: (0.7979, 2.659, 0, 3.267), 4: (0.9213, 1.628, 0, 2.266), 5: (0.9400, 1.427, 0, 2.089),
             6: (0.9515, 1.287, 0.030, 1.970), 10: (0.9727, 0.975, 0.284, 1.716), 25: (0.9896, 0.606, 0.565, 1.435)}

    def test_against_published_tables(self):
        for n, (c4, a3, b3, b4) in self.TABLE.items():
            k = engine.xbar_s_constants(n)
            with self.subTest(n=n):
                self.assertAlmostEqual(k["c4"], c4, delta=0.0006)
                self.assertAlmostEqual(k["A3"], a3, delta=0.0015)
                self.assertAlmostEqual(k["B3"], b3, delta=0.0015)
                self.assertAlmostEqual(k["B4"], b4, delta=0.0015)


def builds_from(groups):
    rows = []
    for i, g in enumerate(groups):
        for j, v in enumerate(g):
            rows.append({"id": f"B{i}", "machine": "M", "coupon": str(j + 1), "strength": v})
    return engine.build_level(rows)


class XbarSChart(unittest.TestCase):
    def test_equal_n_by_hand(self):
        groups = [[1, 2, 3], [2, 3, 4], [3, 4, 5]]            # means 2, 3, 4; each s = 1
        xs = engine.xbar_s(builds_from(groups), baseline=3)
        sigma = 1 / engine.c4(3)
        self.assertAlmostEqual(xs.center, 3.0)
        self.assertAlmostEqual(xs.s_bar, 1.0)
        self.assertAlmostEqual(xs.sigma, sigma)
        self.assertAlmostEqual(xs.x_ucl[0], 3 + 3 * sigma / math.sqrt(3))
        self.assertAlmostEqual(xs.x_ucl[0] - 3, engine.xbar_s_constants(3)["A3"] * 1.0)   # = A3·s̄
        self.assertAlmostEqual(xs.s_ucl[0], engine.xbar_s_constants(3)["B4"] * 1.0)       # = B4·s̄
        self.assertEqual(xs.s_lcl[0], 0.0)
        self.assertEqual(xs.method, "s̄ / c4(3)")

    def test_unequal_n_uses_pooled_sd(self):
        groups = [[10, 12], [11, 13, 15], [9, 11, 13, 15]]
        xs = engine.xbar_s(builds_from(groups), baseline=3)
        dof = 1 + 2 + 3
        sp = math.sqrt((1 * 2.0 + 2 * 4.0 + 3 * (20 / 3)) / dof)
        self.assertAlmostEqual(xs.s_bar, sp)
        self.assertAlmostEqual(xs.sigma, sp / engine.c4(dof + 1))
        self.assertTrue(xs.method.startswith("pooled"))
        self.assertNotEqual(xs.x_ucl[0], xs.x_ucl[2])          # limits step with n

    def test_s_chart_flags_growing_spread(self):
        groups = [[100 + d, 101 - d, 100, 101] for d in (0.5, 1, 0.6, 0.9, 0.7, 1.1, 0.8, 0.6, 1, 0.7)] + [[90, 110, 100, 101]]
        xs = engine.xbar_s(builds_from(groups), baseline=10)
        self.assertEqual(xs.s_hits, {10: "above"})

    def test_build_level_aggregation(self):
        rows = [{"id": "B1", "machine": "M", "coupon": "a", "strength": 950.0, "oxygen": 60.0},
                {"id": "B1", "machine": "M", "coupon": "b", "strength": 960.0, "oxygen": None},
                {"id": "B2", "machine": "M", "coupon": "a", "strength": 940.0, "oxygen": 70.0}]
        b = engine.build_level(rows)
        self.assertEqual([r["id"] for r in b], ["B1", "B2"])
        self.assertEqual((b[0]["strength"], b[0]["n_coupons"], b[0]["strength_min"], b[0]["oxygen"]), (955.0, 2, 950.0, 60.0))
        plain = [{"id": "B1", "machine": "M", "strength": 950.0}]
        self.assertIs(engine.build_level(plain), plain)          # no coupon column → unchanged

    def test_build_fails_if_any_coupon_out_of_spec(self):
        a = engine.analyze_stream(COUPONS, usl=1100)
        self.assertEqual([r["id"] for r in a.fails], ["E-037", "E-040"])
        self.assertTrue(all(r["strength"] > 900 for r in a.fails))   # build means pass; a coupon fails


class CouponImport(unittest.TestCase):
    def test_parse_and_roundtrip(self):
        self.assertEqual(len(COUPONS), 160)
        self.assertEqual(COUPONS[0]["coupon"], "P1")
        self.assertEqual(data.parse_csv(data.to_csv(COUPONS).encode()), COUPONS)

    def test_duplicate_rules(self):
        text = "id,machine,strength\n" + "\n".join(f"B{i // 2},M,{950 + i % 7}" for i in range(60))
        with self.assertRaises(data.ImportError_) as cm:
            data.parse_csv(text.encode())
        self.assertIn("coupon", str(cm.exception))
        dup = "id,machine,coupon,strength\n" + "\n".join(f"B{i},M,1,{950 + i % 7}" for i in range(25)) + "\nB0,M,1,951"
        with self.assertRaises(data.ImportError_):
            data.parse_csv(dup.encode())

    def test_minimum_builds_for_subgroups(self):
        few = "id,machine,coupon,strength\n" + "\n".join(f"B{i},M,{c},{950 + (i * 3 + c) % 11}" for i in range(15) for c in (1, 2, 3))
        with self.assertRaises(data.ImportError_) as cm:
            data.parse_csv(few.encode())
        self.assertIn("at least 20", str(cm.exception))
        ok = "id,machine,coupon,strength\n" + "\n".join(f"B{i},M,{c},{950 + (i * 3 + c) % 11}" for i in range(20) for c in (1, 2, 3))
        self.assertEqual(len(data.parse_csv(ok.encode())), 60)


class CouponDatasetStory(unittest.TestCase):
    a = engine.analyze_stream(COUPONS, usl=1100)

    def test_chart_type_and_limits(self):
        a = self.a
        self.assertEqual(a.chart, "Xbar-S")
        self.assertEqual(a.n, 40)
        self.assertEqual(len(a.values), 160)
        self.assertEqual(a.xs.n, [4] * 40)
        self.assertAlmostEqual(a.xs.center, 967.8, delta=0.05)
        self.assertAlmostEqual(a.xs.s_ucl[0], 16.60, delta=0.01)

    def test_s_chart_signals_first(self):
        self.assertEqual(min(self.a.xs.s_hits) + 1, 30)
        self.assertEqual(min(self.a.xs.x_hits) + 1, 34)
        self.assertTrue(all(f.index + 1 >= 30 for f in self.a.flags))
        self.assertIn("S", self.a.test_counts())

    def test_imr_on_build_means_would_be_later(self):
        means_only = [dict(r, coupon=None) for r in engine.build_level(COUPONS)]
        imr = engine.analyze_stream(means_only, usl=1100)
        self.assertEqual(imr.chart, "I-MR")
        self.assertEqual(imr.flags[0].index + 1, 34)

    def test_not_normal_and_ppm_underestimates(self):
        a = self.a
        self.assertFalse(a.normal["normal"])
        observed_ppm = sum(v < 900 for v in a.values) / len(a.values) * 1e6
        self.assertLess(a.expected["overall"]["total_ppm"], observed_ppm / 100)

    def test_capability_uses_within_build_sigma(self):
        self.assertEqual(self.a.cap_basis, "within-build σ̂ = s̄ / c4(4)")
        self.assertAlmostEqual(self.a.cap["Cpl"], (self.a.avg - 900) / (3 * self.a.xs.sigma))


class CouponViewsAndExports(unittest.TestCase):
    def ctx(self):
        return exports.build_context(COUPONS, dict(engine.DEFAULT_CONFIG), machine="LPBF-E", source="sample:coupons",
                                     actions={}, tasks=views.TASKS, generated_at="2026-10-23 12:00")

    def test_pages(self):
        a = engine.analyze_stream(COUPONS, usl=1100)
        desc = engine.describe_all(COUPONS)
        self.assertEqual(desc["strength"]["n"], 160)
        self.assertEqual(desc["oxygen"]["n"], 40)
        pc = views.render_page("Process control", a, a, dict(engine.DEFAULT_CONFIG), {}, desc)
        self.assertIn("X̄ chart", pc)
        self.assertIn("S chart", pc)
        cap = views.render_page("Capability", a, a, dict(engine.DEFAULT_CONFIG), {}, desc)
        self.assertIn("not normally distributed", cap)
        self.assertIn("Normal probability plot", cap)
        for p in views.PAGES:
            self.assertTrue(views.render_page(p, a, a, dict(engine.DEFAULT_CONFIG), {}, desc))
        self.assertIn("P4: ", views.build_detail(a, 39))

    def test_html_and_json(self):
        ctx = self.ctx()
        html = exports.html_report(ctx)
        self.assertIn("X̄-S chart", html)
        self.assertIn("Distribution check", html)
        rep = json.loads(exports.json_report(ctx))
        self.assertEqual(rep["control_chart"]["type"], "Xbar-S")
        self.assertFalse(rep["normality"]["normal"])

    @unittest.skipUnless(openpyxl, "openpyxl not installed")
    def test_excel(self):
        wb = openpyxl.load_workbook(io.BytesIO(exports.excel_report(self.ctx())))
        self.assertIn("Coupons", wb.sheetnames)
        self.assertEqual(wb["Coupons"].max_row, 161)
        caps = {r[0]: r[1] for r in wb["Capability"].iter_rows(values_only=True) if r[0]}
        self.assertEqual(caps["Normal (p ≥ 0.05)"], "no")

    def test_cli(self):
        out = subprocess.run([sys.executable, str(ROOT / "python" / "analyze.py"), str(ROOT / "datasets" / "04_coupon_spread.csv"),
                              "--machine", "LPBF-E", "--usl", "1100"], capture_output=True, text=True, check=True).stdout
        rep = json.loads(out)
        self.assertEqual(rep["chart"], "Xbar-S")
        self.assertEqual(rep["n"], 40)                      # builds
        self.assertEqual(sum(rep["coupons_per_build"]), 160)
        self.assertEqual(rep["failed_strength_tests"], 2)


if __name__ == "__main__":
    unittest.main()
