"""Excel import and report exports (HTML, Excel, JSON). Run: python -m unittest discover tests

Excel tests are skipped automatically if openpyxl isn't installed (pip install -r requirements.txt).
"""
import datetime as dt
import io
import json
import pathlib
import subprocess
import sys
import unittest

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from forge_sentinel import data, engine, exports, views  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

needs_openpyxl = unittest.skipUnless(openpyxl, "openpyxl not installed")
DRIFT = data.parse_csv((ROOT / "datasets" / "02_oxygen_drift.csv").read_bytes())
MINIMAL = data.parse_csv((ROOT / "datasets" / "03_strength_only.csv").read_bytes())
ACTIONS = {"gas": {"status": "Investigating", "note": "Door seal <worn>", "updated": "2026-10-23 10:00"}}


def ctx_for(rows, source="sample:drift"):
    m = data.machine_list(rows)[0]
    return exports.build_context([r for r in rows if r["machine"] == m], dict(engine.DEFAULT_CONFIG), machine=m,
                                 source=source, actions=ACTIONS, tasks=views.TASKS, generated_at="2026-10-23 11:00")


@needs_openpyxl
class ExcelImport(unittest.TestCase):
    def test_roundtrip_matches_csv(self):
        self.assertEqual(data.read_upload(data.to_xlsx(DRIFT), "builds.xlsx"), DRIFT)

    def test_shipped_excel_dataset_matches_csv(self):
        raw = (ROOT / "datasets" / "02_oxygen_drift.xlsx").read_bytes()
        self.assertEqual(data.read_upload(raw, "02_oxygen_drift.xlsx"), DRIFT)

    def test_real_world_workbook(self):
        """Cover sheet first, Title-Case headers, a date column, blank cells, IDs stored as numbers."""
        wb = openpyxl.Workbook()
        wb.active.title = "Cover"
        wb.active["A1"] = "Quality records"
        ws = wb.create_sheet("Builds")
        ws.append(["ID", "Machine", "Build date", "Strength", "Oxygen"])
        for i in range(35):
            ws.append([float(i + 1), "M1", dt.date(2026, 7, 1) + dt.timedelta(days=i), 950 + (i * 7) % 23,
                       None if i % 4 == 0 else 60.5])
        buf = io.BytesIO()
        wb.save(buf)
        rows = data.read_upload(buf.getvalue(), "Records.XLSX")
        self.assertEqual(len(rows), 35)
        self.assertEqual(rows[0]["id"], "1")
        self.assertEqual(sum(r["oxygen"] is None for r in rows), 9)

    def test_rejections(self):
        for raw, name, msg in [(b"not a workbook", "a.xlsx", "Could not open"), (b"x", "old.xls", ".xls")]:
            with self.subTest(name), self.assertRaises(data.ImportError_) as cm:
                data.read_upload(raw, name)
            self.assertIn(msg, str(cm.exception))

    def test_cli_reads_excel(self):
        tmp = ROOT / "tests" / "_tmp.xlsx"
        tmp.write_bytes(data.to_xlsx(DRIFT))
        try:
            out = subprocess.run([sys.executable, str(ROOT / "python" / "analyze.py"), str(tmp), "--machine", "LPBF-C",
                                  "--usl", "1100"], capture_output=True, text=True, check=True).stdout
        finally:
            tmp.unlink()
        self.assertEqual(json.loads(out)["failed_strength_tests"], 3)


class HtmlReport(unittest.TestCase):
    def test_contents(self):
        html = exports.html_report(ctx_for(DRIFT))
        for needle in ("PROCESS ASSURANCE REPORT", "Process unstable", "Print / Save as PDF", "<svg",
                       "Test dataset: Oxygen drift", "C-056", "Chamber oxygen", "Door seal &lt;worn&gt;", "Wilson"):
            self.assertIn(needle, html)
        self.assertNotIn("nan", html.lower().replace("nancy", ""))

    def test_oxygen_ranked_first_in_early_warnings(self):
        w = exports.early_warnings(ctx_for(DRIFT))
        self.assertEqual(w[0]["v"], "oxygen")
        self.assertEqual(w[0]["first_seq"], 44)
        self.assertEqual(w[0]["before_strength"], 5)

    def test_strength_only(self):
        html = exports.html_report(ctx_for(MINIMAL, "sample:minimal"))
        self.assertIn("Only strength was recorded", html)


class JsonReport(unittest.TestCase):
    def test_structure(self):
        rep = json.loads(exports.json_report(ctx_for(DRIFT)))
        self.assertEqual(rep["conformance"]["failed_builds"], ["C-056", "C-057", "C-060"])
        self.assertEqual(rep["process_variable_signals"]["oxygen"][0]["build"], "C-044")
        self.assertIn("capability", rep)


@needs_openpyxl
class ExcelReport(unittest.TestCase):
    def test_workbook(self):
        wb = openpyxl.load_workbook(io.BytesIO(exports.excel_report(ctx_for(DRIFT))))
        self.assertEqual(wb.sheetnames, ["Summary", "Data", "Statistics", "Capability", "Signals", "Associations",
                                         "Corrective actions"])
        self.assertEqual(len(wb["Summary"]._charts), 1)
        data_ws = wb["Data"]
        self.assertEqual(data_ws.max_row, 71)                       # header + 70 builds
        headers = [c.value for c in data_ws[1]]
        results = [row[headers.index("Result")] for row in data_ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(results.count("FAIL"), 3)
        summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
        self.assertEqual(summary["Verdict"], "Process unstable — capability not demonstrated")

    def test_strength_only_workbook(self):
        wb = openpyxl.load_workbook(io.BytesIO(exports.excel_report(ctx_for(MINIMAL, "sample:minimal"))))
        self.assertEqual(wb["Associations"]["A2"].value, "(only strength recorded)")


if __name__ == "__main__":
    unittest.main()
