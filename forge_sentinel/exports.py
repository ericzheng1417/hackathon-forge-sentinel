"""Report exports: printable HTML (save as PDF from the browser), Excel workbook, and JSON.

All three are built from one `ReportContext`, so they always agree with each other and
with the dashboard. HTML and JSON use only the standard library; the Excel workbook
needs `openpyxl` (imported inside `excel_report`).
"""
from __future__ import annotations

import io
import json
import math
from dataclasses import dataclass, field
from html import escape as esc

from . import charts, engine
from .charts import fmt, num
from .engine import SPECIAL_TESTS, VARS

STATUSES_DEFAULT = "Open"


@dataclass
class ReportContext:
    """Everything a report needs for ONE machine stream."""
    machine: str
    source: str
    generated_at: str
    cfg: dict
    actions: dict
    tasks: list                      # (id, title, description, owner) from views.TASKS
    a: engine.Analysis               # strength analysis
    process: dict = field(default_factory=dict)   # variable -> Analysis for recorded process variables
    desc: dict = field(default_factory=dict)      # variable -> describe()
    assoc: list = field(default_factory=list)
    wilson: tuple = (math.nan, math.nan, math.nan)

    @property
    def source_label(self) -> str:
        if self.source == "demo":
            return "Built-in demonstration data"
        if self.source.startswith("sample:"):
            from .data import SAMPLE_DATASETS
            fname, label = SAMPLE_DATASETS.get(self.source[7:], (self.source[7:], self.source[7:]))
            return f"Test dataset: {label} ({fname})"
        return "Imported file"


def build_context(stream: list[dict], cfg: dict, *, machine: str, source: str, actions: dict, tasks: list,
                  generated_at: str) -> ReportContext:
    """Run every analysis the reports need, with the dashboard's current settings."""
    kw = dict(lsl=cfg["lsl"], baseline=int(cfg["baseline"]), target=cfg["target"], usl=cfg.get("usl") or None,
              tests=cfg.get("tests", engine.ALL_TESTS), mr_rule=cfg.get("mr_rule", True))
    a = engine.analyze_stream(stream, "strength", **kw)
    builds = engine.build_level(stream)
    avail = engine.available_vars(builds)
    process = {v: engine.analyze_stream(stream, v, **kw) for v in avail if v != "strength"}
    return ReportContext(
        machine=machine, source=source, generated_at=generated_at, cfg=cfg, actions=actions, tasks=tasks, a=a,
        process=process, desc=engine.describe_all(stream),
        assoc=engine.associations(builds), wilson=engine.wilson(a.passes, a.n))


def json_report(ctx: ReportContext) -> str:
    """Machine-readable export (same structure as before, plus process-variable signals)."""
    rep = engine.report(ctx.a, source=ctx.source, machine=ctx.machine, config=ctx.cfg, actions=ctx.actions,
                        generated_at=ctx.generated_at)
    rep["process_variable_signals"] = engine._clean({
        v: [{"build": am.records[f.index]["id"], "tests": f.tests, "reasons": f.reasons} for f in am.flags]
        for v, am in ctx.process.items()})
    p, lo, hi = ctx.wilson
    rep["conformance"] = engine._clean({"pass_rate": p, "wilson_95": [lo, hi], "failed_builds": [r["id"] for r in ctx.a.fails]})
    return json.dumps(rep, indent=2, ensure_ascii=False)


# -- shared summaries ------------------------------------------------------

def key_metrics(ctx: ReportContext) -> list[tuple[str, str, str]]:
    """(label, value, note) rows used by both the HTML and Excel summaries."""
    a, c, d = ctx.a, ctx.a.cap, ctx.a.desc
    p, lo, hi = ctx.wilson
    t = ctx.cfg["target"]
    return [
        ("Verdict", engine.decision_text(a), "Signals → unstable; otherwise Cpk is compared with the target"),
        ("Builds analysed", str(a.n), f"Baseline: first {a.baseline} builds"),
        ("Strength signals (monitoring)", str(len(a.post)), f"{len(a.baseflags)} in baseline"),
        ("Out-of-spec builds", str(len(a.fails)), ", ".join(str(r["id"]) for r in a.fails) or "none"),
        ("Observed pass rate", f"{fmt(p * 100, 1)}%", f"95% Wilson interval {fmt(lo * 100, 1)}–{fmt(hi * 100, 1)}%"),
        ("Cpk", fmt(c["Cpk"]), f"target ≥ {t} · {'meets' if c['Cpk'] >= t else 'below'} target"),
        ("Ppk", fmt(c["Ppk"]), "overall spread (includes any drift)"),
        ("Cp / Pp", f"{fmt(c['Cp'])} / {fmt(c['Pp'])}", "potential (spread only)" if a.usl is not None else "need USL"),
        ("Mean / median strength", f"{fmt(d['mean'], 1)} / {fmt(d['median'], 1)} MPa", ""),
        ("SD / CV", f"{fmt(d['sd'], 2)} MPa / {fmt(d['cv_pct'], 2)}%", f"within σ {fmt(a.sigma, 2)} MPa ({a.cap_basis})"),
        ("Control chart", "X̄-S (coupons per build)" if a.xs else "I-MR (one result per build)",
         f"{len(a.values)} individual results" if a.xs else ""),
        ("Normality (Anderson-Darling)", _pval(a.normal), "p ≥ 0.05: consistent with normal; p < 0.05: not normal"),
        ("Expected out of spec (normal fit)", f"{_ppm(a.expected.get('overall', {}).get('total_ppm'))} ppm (overall s)",
         f"{_ppm(a.expected.get('within', {}).get('total_ppm'))} ppm using within σ"),
    ]


def _pval(nm: dict) -> str:
    p = (nm or {}).get("p")
    if not isinstance(p, float) or not math.isfinite(p):
        return "—"
    return ("p < 0.005" if p < 0.005 else f"p = {p:.3f}") + (" (normal)" if nm["normal"] else " (not normal)")


def _ppm(v) -> str:
    from .views import ppm
    return ppm(v)


def early_warnings(ctx: ReportContext) -> list[dict]:
    """Process variables with signals, ranked by |correlation with strength|, then by first signal."""
    r_by_v = {x["v"]: x["r"] for x in ctx.assoc}
    out = []
    for v, am in ctx.process.items():
        if am.flags:
            f = am.flags[0]
            out.append({"v": v, "label": VARS[v][0], "count": len(am.flags), "first_build": am.records[f.index]["id"],
                        "first_seq": am.idx[f.index] + 1, "reasons": f.reasons, "r": r_by_v.get(v, math.nan)})
    first_strength = ctx.a.flags[0].index + 1 if ctx.a.flags else None
    for w in out:
        w["before_strength"] = (first_strength - w["first_seq"]) if first_strength else None
    # rank by strength of association: isolated signals on unrelated variables are usually false alarms
    return sorted(out, key=lambda w: (-abs(w["r"]) if math.isfinite(w["r"]) else 0, w["first_seq"]))


def _spec(a: engine.Analysis) -> str:
    return f"{num(a.lsl)}–{num(a.usl)} MPa" if a.usl is not None else f"≥ {num(a.lsl)} MPa (one-sided)"


# -- printable HTML ---------------------------------------------------------

REPORT_CSS = """
:root{--ink:#202c32;--muted:#6b7a7f;--green:#247568;--line:#dfe6e3;--red:#bd594a;--amber:#9e7734}
*{box-sizing:border-box}
html,body{margin:0;background:#f3f5f4;color:var(--ink);font:13px/1.55 "DM Sans",Arial,Helvetica,sans-serif}
.page{max-width:900px;margin:24px auto;background:#fff;padding:40px 48px;border:1px solid var(--line);border-radius:8px}
h1{font:800 24px Manrope,Arial,sans-serif;margin:0 0 4px}
h2{font:700 16px Manrope,Arial,sans-serif;margin:30px 0 10px;padding-bottom:6px;border-bottom:2px solid var(--line)}
h3{font-size:13px;margin:16px 0 6px}
p{margin:6px 0;color:#41505a}
.muted{color:var(--muted)}
.meta{display:grid;grid-template-columns:repeat(3,1fr);gap:4px 18px;font-size:12px;color:var(--muted);margin:10px 0 18px}
.meta b{color:var(--ink);font-weight:600}
.verdict{border-radius:8px;padding:14px 18px;margin:10px 0 18px;border:1px solid}
.verdict.bad{background:#fcf1ee;border-color:#ecc9bf;color:#8a3b2d}
.verdict.warn{background:#fcf6ea;border-color:#ead8b0;color:#7a5a1e}
.verdict.ok{background:#eef6f2;border-color:#c5dfd3;color:#1e5e4f}
.verdict strong{font-size:15px;display:block}
.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:8px 0}
.kpi{border:1px solid var(--line);border-radius:6px;padding:10px 12px}
.kpi .l{font-size:10px;letter-spacing:.4px;text-transform:uppercase;color:var(--muted)}
.kpi .v{font:700 20px Manrope,Arial,sans-serif;margin-top:2px}
.kpi .n{font-size:10.5px;color:var(--muted)}
table{border-collapse:collapse;width:100%;font-size:11.5px;margin:6px 0 10px}
th{background:#f4f7f5;color:var(--muted);font-weight:600;font-size:10px;letter-spacing:.4px;text-transform:uppercase;text-align:left;padding:7px 8px}
td{padding:6px 8px;border-bottom:1px solid #eef2f0;vertical-align:top}
td.num,th.num{text-align:right;font-variant-numeric:tabular-nums}
tr.off td{color:#b3bcb8}
.pill{display:inline-block;font-size:10px;border-radius:4px;padding:1px 6px;background:#eef3f0;color:#427568}
.pill.red{background:#f9e9e5;color:var(--red)}
.chart{width:100%;display:block;margin:4px 0}
.chart text{font:10px Arial,Helvetica,sans-serif;fill:#7d8a8a}
.chart text.zone{fill:#b7c4be;font-weight:700;font-size:9px}
.legend{display:flex;gap:16px;font-size:10px;color:var(--muted);margin:6px 0}
.legend span:before{content:"";display:inline-block;width:13px;height:2px;background:var(--green);vertical-align:middle;margin-right:5px}
.legend .control:before{background:#b2bab7}.legend .spec:before{background:var(--red)}
.legend .zones:before{background:repeating-linear-gradient(90deg,#c9d6d0 0 2px,transparent 2px 5px)}
.two{display:grid;grid-template-columns:1.2fr 1fr;gap:20px;align-items:start}
.toolbar{max-width:900px;margin:16px auto 0;display:flex;gap:10px;justify-content:flex-end}
.toolbar button{font:inherit;border:1px solid var(--green);background:var(--green);color:#fff;border-radius:6px;padding:8px 14px;cursor:pointer}
.note{font-size:11px;color:var(--muted)}
@page{size:A4;margin:14mm}
@media print{
  html,body{background:#fff}
  .toolbar{display:none}
  .page{margin:0;border:0;padding:0;max-width:none}
  h2,h3{break-after:avoid}
  .chart,.kpis,.verdict,.two{break-inside:avoid}
  tr{break-inside:avoid}
}
@media (max-width:700px){.page{padding:20px 16px;margin:0}.kpis{grid-template-columns:repeat(2,1fr)}.meta,.two{grid-template-columns:1fr}}
"""


def html_report(ctx: ReportContext) -> str:
    """Self-contained, printable HTML report (open in a browser → Print → Save as PDF)."""
    a, c, cfg = ctx.a, ctx.a.cap, ctx.cfg
    verdict = engine.decision_text(a)
    vcls = "bad" if not a.stable else ("warn" if not c["Cpk"] >= cfg["target"] else "ok")
    p, lo, hi = ctx.wilson
    tests_on = ", ".join(map(str, a.tests)) or "none"
    counts = a.test_counts()

    kpis = "".join(f'<div class="kpi"><div class="l">{esc(l)}</div><div class="v">{v}</div><div class="n">{esc(n)}</div></div>'
                   for l, v, n in [
                       ("Pass rate", f"{fmt(p * 100, 1)}%", f"{a.passes}/{a.n} · 95% CI {fmt(lo * 100, 1)}–{fmt(hi * 100, 1)}%"),
                       ("Cpk", fmt(c["Cpk"]), f"target ≥ {num(cfg['target'])}"),
                       ("Ppk", fmt(c["Ppk"]), "overall spread"),
                       ("Strength signals", str(len(a.post)), f"after baseline · {len(a.baseflags)} in baseline"),
                       ("Mean strength", f"{fmt(a.desc['mean'], 1)}", f"MPa · median {fmt(a.desc['median'], 1)}"),
                       ("Std deviation", f"{fmt(a.desc['sd'], 2)}", f"MPa · CV {fmt(a.desc['cv_pct'], 2)}%"),
                       ("Cp / Pp", f"{fmt(c['Cp'])} / {fmt(c['Pp'])}", "potential (spread only)"),
                       ("Out of spec", str(len(a.fails)), _spec(a)),
                   ])

    warnings = early_warnings(ctx)
    if warnings:
        warn_rows = "".join(
            f'<tr><td>{esc(w["label"])}</td><td class="num">{w["count"]}</td><td>{esc(str(w["first_build"]))} (build {w["first_seq"]})</td>'
            f'<td>{esc("; ".join(w["reasons"]))}</td><td class="num">{fmt(w["r"])}</td>'
            f'<td>{"" if w["before_strength"] is None else (str(w["before_strength"]) + " builds earlier" if w["before_strength"] > 0 else ("same build" if w["before_strength"] == 0 else str(-w["before_strength"]) + " builds later"))}</td></tr>'
            for w in warnings)
        warn_html = ('<table><thead><tr><th>Variable</th><th class="num">Signals</th><th>First flagged</th><th>Reason</th>'
                     f'<th class="num">r with strength</th><th>vs first strength signal</th></tr></thead><tbody>{warn_rows}</tbody></table>')
        warn_html += ('<p class="note">Ranked by association with strength. A variable that is both flagged and strongly '
                      'associated is the first place to investigate; isolated signals on weakly associated variables are '
                      'often false alarms (7 variables × 8 tests give many chances).</p>')
    elif ctx.process:
        warn_html = '<p class="muted">No signals on the recorded process variables.</p>'
    else:
        warn_html = '<p class="muted">Only strength was recorded; add process columns to find causes.</p>'

    if ctx.assoc:
        top = ctx.assoc[0]
        assoc_rows = "".join(f'<tr><td>{esc(VARS[x["v"]][0])}</td><td class="num">{fmt(x["r"])}</td><td class="num">{x["n"]}</td></tr>'
                             for x in ctx.assoc)
        assoc_html = (f'<table><thead><tr><th>Variable</th><th class="num">Pearson r</th><th class="num">Pairs</th></tr></thead>'
                      f'<tbody>{assoc_rows}</tbody></table>'
                      f'<p class="note">Strongest: <b>{esc(VARS[top["v"]][0].lower())}</b> (r = {fmt(top["r"])}). '
                      'Correlation suggests where to look; it does not prove cause.</p>')
        scatter = charts.scatter(a, a.lsl, top["v"])
    else:
        assoc_html, scatter = '<p class="muted">No process variables recorded.</p>', ""

    test_rows = "".join(
        f'<tr class="{"" if t in a.tests else "off"}"><td><b>Test {t}</b></td><td>{esc(pat)}</td>'
        f'<td class="num">{(counts.get(t, 0) or "0") if t in a.tests else "off"}</td></tr>'
        for t, (_, pat, _sig) in SPECIAL_TESTS.items())
    if a.xs:
        test_rows += (f'<tr><td><b>S</b></td><td>Within-build SD outside the S-chart limits</td>'
                      f'<td class="num">{counts.get("S", 0) or "0"}</td></tr>')
    else:
        test_rows += (f'<tr class="{"" if cfg.get("mr_rule", True) else "off"}"><td><b>MR</b></td><td>Moving range above the MR-chart UCL</td>'
                      f'<td class="num">{(counts.get("MR", 0) or "0") if cfg.get("mr_rule", True) else "off"}</td></tr>')

    if a.flags:
        sig_rows = "".join(
            f'<tr><td>{esc(str(a.records[f.index]["id"]))}</td><td class="num">{f.index + 1}</td>'
            f'<td>{"baseline" if f.index < a.baseline else "monitoring"}</td><td class="num">{num(a.records[f.index]["strength"])}</td>'
            f'<td>{esc(", ".join(map(str, f.tests)))}</td></tr>' for f in a.flags)
        sig_html = ('<table><thead><tr><th>Build</th><th class="num">Seq.</th><th>Phase</th><th class="num">Strength (MPa)</th>'
                    f'<th>Tests failed</th></tr></thead><tbody>{sig_rows}</tbody></table>')
    else:
        sig_html = '<p class="muted">No strength signals under the selected tests.</p>'

    fail_html = ("".join(f'<tr><td>{esc(str(r["id"]))}</td><td class="num">{num(r["strength"])}'
                         + (f' <span class="muted">(min {fmt(r["strength_min"], 1)})</span>' if r.get("n_coupons", 1) > 1 else "")
                         + '</td>'
                         + "".join(f'<td class="num">{num(r.get(v))}</td>' for v in ctx.process)
                         + "</tr>" for r in a.fails))
    if a.fails:
        fail_html = ('<table><thead><tr><th>Build</th><th class="num">Strength</th>'
                     + "".join(f'<th class="num">{esc(VARS[v][0])}</th>' for v in ctx.process)
                     + f'</tr></thead><tbody>{fail_html}</tbody></table>')
    else:
        fail_html = '<p class="muted">No out-of-specification builds.</p>'

    stat_cols = [("n", "N", 0), ("mean", "Mean", 2), ("median", "Median", 2), ("sd", "SD", 2), ("variance", "Var", 2),
                 ("cv_pct", "CV %", 2), ("se_mean", "SE", 3), ("min", "Min", 1), ("max", "Max", 1), ("range", "Range", 1),
                 ("iqr", "IQR", 2)]
    desc_rows = "".join(f'<tr><td>{esc(VARS[v][0])} <span class="muted">{VARS[v][1]}</span></td>'
                        + "".join(f'<td class="num">{fmt(d[k], dec)}</td>' for k, _, dec in stat_cols) + "</tr>"
                        for v, d in ctx.desc.items())
    desc_html = ('<table><thead><tr><th>Variable</th>' + "".join(f'<th class="num">{h}</th>' for _, h, _ in stat_cols)
                 + f'</tr></thead><tbody>{desc_rows}</tbody></table>')

    cap_rows = "".join(f'<tr><td>{k}</td><td class="num">{fmt(c[k])}</td><td>{note}</td></tr>' for k, note in [
        ("Cp", "(USL − LSL) / 6σ within"), ("Cpk", "min(Cpl, Cpu), within σ"), ("Cpl", "(mean − LSL) / 3σ"),
        ("Cpu", "(USL − mean) / 3σ"), ("Pp", "(USL − LSL) / 6s overall"), ("Ppk", "min(Ppl, Ppu), overall s")])

    nm = a.normal or {}
    exp_w, exp_o = a.expected.get("within", {}), a.expected.get("overall", {})
    from .views import dist_curves
    if a.xs:
        chart_html = (f'{charts.xbar_chart(a)}<p class="note">Grand mean {fmt(a.xs.center, 1)} MPa · within-build σ̂ '
                      f'({esc(a.xs.method)}) {fmt(a.xs.sigma, 2)} · X̄ limits x̿ ± 3σ̂/√n. Shaded = baseline. Red = signals.</p>'
                      f'{charts.s_chart(a)}<p class="note">S chart: S̄ {fmt(a.xs.s_bar, 2)} · UCL {fmt(a.xs.s_ucl[0], 2)} · '
                      f'LCL {fmt(a.xs.s_lcl[0], 2)}. A rising S means coupons on the same plate are drifting apart.</p>')
        chart_title = "X̄-S chart"
    else:
        chart_html = (f'{charts.control_chart(a, a.lsl, a.usl)}'
                      f'<p class="note">Centre {fmt(a.center, 1)} MPa · σ (MR̄/1.128) {fmt(a.sigma, 2)} · LCL {fmt(a.lcl, 1)} · '
                      f'UCL {fmt(a.ucl, 1)}. Shaded area = baseline. Red points = special-cause signals.</p>'
                      f'{charts.mr_chart(a)}<p class="note">Moving range: MR̄ {fmt(a.mr_bar, 2)} · UCL {fmt(a.mr_ucl, 1)}.</p>')
        chart_title = "I-MR chart"
    norm_html = (f'<table><tbody><tr><td>Anderson-Darling A² / A²*</td><td class="num">{fmt(nm.get("ad"), 3)} / {fmt(nm.get("ad_adj"), 3)}</td></tr>'
                 f'<tr><td>p-value</td><td class="num">{_pval(nm)}</td></tr>'
                 f'<tr><td>Expected out of spec, within σ</td><td class="num">{_ppm(exp_w.get("total_ppm"))} ppm</td></tr>'
                 f'<tr><td>Expected out of spec, overall s</td><td class="num">{_ppm(exp_o.get("total_ppm"))} ppm</td></tr>'
                 f'<tr><td>Observed out of spec</td><td class="num">{sum(1 for v in a.values if v < a.lsl or (a.usl is not None and v > a.usl))}'
                 f' of {len(a.values)} {"coupons" if a.xs else "builds"}</td></tr></tbody></table>'
                 + ('<p class="note"><b>Not normal:</b> treat the indices and ppm estimates with caution; rely on the observed '
                    'failures and the charts.</p>' if nm.get("normal") is False else ""))

    action_rows = "".join(
        f'<tr><td><b>{esc(title)}</b><div class="note">{esc(owner)}</div></td>'
        f'<td>{esc(ctx.actions.get(tid, {}).get("status", STATUSES_DEFAULT))}</td>'
        f'<td>{esc(ctx.actions.get(tid, {}).get("note", "")) or "<span class=muted>—</span>"}</td>'
        f'<td class="note">{esc(ctx.actions.get(tid, {}).get("updated", ""))}</td></tr>'
        for tid, title, _desc, owner in ctx.tasks)

    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Forge Sentinel report · {esc(ctx.machine)}</title><style>{REPORT_CSS}</style>
<script>function fsBuild(){{}}</script></head>
<body>
<div class="toolbar"><button onclick="window.print()">Print / Save as PDF</button></div>
<div class="page">
<div class="muted" style="font-size:11px;letter-spacing:1.4px;font-weight:700">FORGE SENTINEL · PROCESS ASSURANCE REPORT</div>
<h1>Ti-6Al-4V bracket strength — machine {esc(ctx.machine)}</h1>
<div class="meta">
<div>Data: <b>{esc(ctx.source_label)}</b></div><div>Builds: <b>{a.n}</b>{(" × " + str(a.xs.typical_n) + " coupons") if a.xs and a.xs.equal_n else ""} (baseline first {a.baseline})</div><div>Generated: <b>{esc(ctx.generated_at)}</b></div>
<div>Specification: <b>{_spec(a)}</b></div><div>Cpk target: <b>{num(cfg['target'])}</b></div><div>Tests: <b>{tests_on}{' + S chart' if a.xs else (' + MR' if cfg.get('mr_rule', True) else '')}</b></div>
</div>
<div class="verdict {vcls}"><strong>{esc(verdict)}</strong>
{len(a.post)} monitored builds with strength signals · {len(a.fails)} builds outside the {_spec(a)} specification ·
{'baseline has signals; review before use' if a.baseflags else 'limits frozen from a signal-free baseline'}.</div>
<div class="kpis">{kpis}</div>

<section><h2>1. Strength stability ({chart_title})</h2>
{chart_html}</section>

<section><h2>2. Tests for special causes</h2>
<table><thead><tr><th>Test</th><th>Pattern</th><th class="num">Points flagged</th></tr></thead><tbody>{test_rows}</tbody></table>
<h3>Signal log (strength)</h3>{sig_html}</section>

<section><h2>3. Capability</h2>
{charts.histogram(a.values, "Strength (MPa)", a.lsl, a.usl, bins=36, width=760, curves=dist_curves(a))}
<div class="two"><div><table><thead><tr><th>Index</th><th class="num">Value</th><th>Definition</th></tr></thead><tbody>{cap_rows}</tbody></table>
<p class="note">Cp/Cpk use within σ = {fmt(a.sigma, 2)} ({esc(a.cap_basis)}); Pp/Ppk use overall s = {fmt(a.overall_sd, 2)} MPa.
Indices are meaningful only for a stable, approximately normal process.</p></div>
<div><h3 style="margin-top:0">Distribution check</h3>{norm_html}</div></div>
{charts.qq_plot(a.values, "Strength", a.normal)}
<h3>Out-of-specification builds</h3>{fail_html}</section>

<section><h2>4. Where to look: process variables</h2>
<h3>Early warnings (process variables with signals)</h3>{warn_html}
<h3>Association with strength</h3>{assoc_html}{scatter}</section>

<section><h2>5. Descriptive statistics</h2>{desc_html}
<p class="note">Sample statistics (n − 1). CV = SD ÷ mean × 100. SE = SD ÷ √n. Quartiles: inclusive method.</p></section>

<section><h2>6. Corrective actions</h2>
<table><thead><tr><th>Action</th><th>Status</th><th>Notes</th><th>Updated</th></tr></thead><tbody>{action_rows}</tbody></table></section>

<section><h2>7. Method and limitations</h2>
<p>{"X̄ and S charts (one subgroup = one build's coupons)" if a.xs else "Individuals and moving-range chart"} with limits frozen from the first {a.baseline} builds; the 8 Western Electric / Nelson
tests {"plus an S-chart limit check" if a.xs else "plus a moving-range rule"}; capability per NIST/SEMATECH definitions; Anderson-Darling normality test
and normal-theory ppm estimates; 95% Wilson score interval for the pass rate;
Pearson correlation on pairwise-complete records. Full formulas: <i>docs/REFERENCE.md</i>.</p>
<p class="note">Decision support only. This report does not certify parts, qualify the process, or estimate fatigue life.
Specification limits and targets shown are the values configured in the tool; confirm them against the contract or drawing.</p>
</section>
</div></body></html>"""


# -- Excel workbook ---------------------------------------------------------

def excel_report(ctx: ReportContext) -> bytes:
    """Formatted Excel workbook: Summary (with a native control chart), Data, Statistics,
    Capability, Signals, Associations, Corrective actions. Needs openpyxl."""
    import openpyxl
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    a, c, cfg = ctx.a, ctx.a.cap, ctx.cfg
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="E9F2EE")
    red_fill, title_font = PatternFill("solid", fgColor="F9E2DC"), Font(bold=True, size=14, color="216957")

    def val(v):
        return None if v is None or (isinstance(v, float) and not math.isfinite(v)) else v

    def table(ws, row, headers, rows, widths=None, numfmt="0.00"):
        for j, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=j, value=h)
            cell.font, cell.fill = bold, head_fill
        for i, r in enumerate(rows, start=row + 1):
            for j, v in enumerate(r, start=1):
                cell = ws.cell(row=i, column=j, value=val(v))
                if isinstance(v, float):
                    cell.number_format = numfmt
        if widths:
            for j, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(j)].width = w
        return row + len(rows) + 2

    wb = openpyxl.Workbook()

    # --- Data sheet (also feeds the chart) ---
    data_ws = wb.active
    data_ws.title = "Data"
    flagged = {f.index: f for f in a.flags}
    pvars = list(ctx.process)
    xs = a.xs
    extra = ["Coupons", "Coupon SD", "Lowest coupon"] if xs else []
    headers = (["Seq", "Build ID", "Machine", "Mean strength (MPa)" if xs else "Strength (MPa)"] + extra
               + [f"{VARS[v][0]} ({VARS[v][1]})" for v in pvars]
               + ["Centre", "UCL", "LCL", "LSL"] + (["USL"] if a.usl is not None else [])
               + (["S", "S UCL"] if xs else ["Moving range"]) + ["Result", "Phase", "Strength signals"])
    rows = []
    for i, r in enumerate(a.records):
        f = flagged.get(i)
        limits = [xs.center, xs.x_ucl[i], xs.x_lcl[i]] if xs else [a.center, a.ucl, a.lcl]
        rows.append([i + 1, r["id"], r["machine"], r["strength"]]
                    + ([r.get("n_coupons"), r.get("strength_sd"), r.get("strength_min")] if xs else [])
                    + [r.get(v) for v in pvars]
                    + limits + [a.lsl] + ([a.usl] if a.usl is not None else [])
                    + ([xs.sds[i], xs.s_ucl[i]] if xs else [a.mr[i]])
                    + ["FAIL" if a.is_fail(r) else "PASS", "baseline" if i < a.baseline else "monitoring",
                       "; ".join(f.reasons) if f else ""])
    table(data_ws, 1, headers, rows, [6, 11, 10, 14] + [10, 11, 13][:len(extra)] + [14] * len(pvars) + [10, 10, 10, 8]
          + ([8] if a.usl is not None else []) + ([9, 9] if xs else [12]) + [8, 11, 48], numfmt="0.0")
    res_col = headers.index("Result") + 1
    for i, r in enumerate(a.records, start=2):
        if a.is_fail(r):
            data_ws.cell(row=i, column=res_col).fill = red_fill
        if (i - 2) in flagged:
            data_ws.cell(row=i, column=4).fill = red_fill
    data_ws.freeze_panes = "C2"
    data_ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    # --- Summary sheet ---
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = f"Forge Sentinel · process assurance report · {ctx.machine}"
    ws["A1"].font = title_font
    meta = [("Data", ctx.source_label), ("Generated", ctx.generated_at), ("Specification", _spec(a)),
            ("Cpk target", cfg["target"]), ("Baseline builds", a.baseline),
            ("Tests applied", ", ".join(map(str, a.tests)) + (" + MR" if cfg.get("mr_rule", True) else ""))]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=k).font = bold
        ws.cell(row=i, column=2, value=v).alignment = Alignment(horizontal="left")
    r0 = 3 + len(meta) + 1
    table(ws, r0, ["Metric", "Value", "Note"], key_metrics(ctx), [30, 46, 52])
    verdict_cell = ws.cell(row=r0 + 1, column=2)
    verdict_cell.font = Font(bold=True, color="A33F2F" if not a.stable else "216957")
    for row in ws.iter_rows(min_row=r0, max_row=r0 + 12, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    chart = LineChart()
    chart.title = "X̄ chart: mean strength per build" if xs else "Strength individuals chart"
    chart.y_axis.title, chart.x_axis.title = "MPa", "Build sequence"
    chart.height, chart.width = 9, 22
    chart.x_axis.delete = False          # recent Excel versions hide axes unless this is explicit
    chart.y_axis.delete = False
    chart.y_axis.number_format = "0"
    chart.x_axis.tickLblSkip = 5
    first_ctrl = 5 + len(extra) + len(pvars)   # Centre, UCL, LCL, LSL columns in the Data sheet
    n_rows = len(rows) + 1
    chart.add_data(Reference(data_ws, min_col=4, min_row=1, max_row=n_rows), titles_from_data=True)
    chart.add_data(Reference(data_ws, min_col=first_ctrl, max_col=first_ctrl + (2 if xs else 3), min_row=1, max_row=n_rows),
                   titles_from_data=True)   # spec limits apply to coupons, so the X̄ chart omits LSL
    chart.set_categories(Reference(data_ws, min_col=1, min_row=2, max_row=n_rows))
    lo = min(*a.x, *(xs.x_lcl if xs else [a.lcl]), *([] if xs else [a.lsl])) - 10
    hi = max(*a.x, *(xs.x_ucl if xs else [a.ucl])) + 10
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = math.floor(lo / 10) * 10, math.ceil(hi / 10) * 10
    colours = ["427E6D", "5B9685", "A1B4AC", "A1B4AC", "C36A57"]
    for s, col in zip(chart.series, colours):
        s.graphicalProperties.line.solidFill = col
        s.graphicalProperties.line.width = 22000 if s is chart.series[0] else 12000
        s.smooth = False
        if s is not chart.series[0]:
            s.graphicalProperties.line.dashStyle = "dash"
    ws.add_chart(chart, f"A{r0 + 14}")   # USL stays off the chart so the scale remains readable

    # --- Statistics ---
    st_ws = wb.create_sheet("Statistics")
    keys = [("n", "N"), ("mean", "Mean"), ("median", "Median"), ("sd", "Std dev"), ("variance", "Variance"),
            ("cv_pct", "CV %"), ("se_mean", "SE mean"), ("min", "Min"), ("q1", "Q1"), ("q3", "Q3"), ("max", "Max"),
            ("range", "Range"), ("iqr", "IQR")]
    table(st_ws, 1, ["Variable", "Unit"] + [h for _, h in keys],
          [[VARS[v][0], VARS[v][1]] + [d[k] for k, _ in keys] for v, d in ctx.desc.items()],
          [24, 8] + [10] * len(keys))
    st_ws.cell(row=len(ctx.desc) + 3, column=1, value="Sample statistics (n − 1). Quartiles: inclusive method (QUARTILE.INC).")

    # --- Capability ---
    cap_ws = wb.create_sheet("Capability")
    nxt = table(cap_ws, 1, ["Index", "Value", "Definition"], [
        ["Cp", c["Cp"], "(USL − LSL) / 6σ within"], ["Cpk", c["Cpk"], "min(Cpl, Cpu)"],
        ["Cpl", c["Cpl"], "(mean − LSL) / 3σ within"], ["Cpu", c["Cpu"], "(USL − mean) / 3σ within"],
        ["Pp", c["Pp"], "(USL − LSL) / 6s overall"], ["Ppk", c["Ppk"], "min(Ppl, Ppu)"],
        ["Ppl", c["Ppl"], "(mean − LSL) / 3s overall"], ["Ppu", c["Ppu"], "(USL − mean) / 3s overall"]], [12, 10, 34])
    p, lo_w, hi_w = ctx.wilson
    nm = a.normal or {}
    exp_w, exp_o = a.expected.get("within", {}), a.expected.get("overall", {})
    table(cap_ws, nxt, ["Quantity", "Value"], [
        ["Mean (MPa)", a.avg], [f"Within σ ({a.cap_basis}) (MPa)", a.sigma], ["Overall s (MPa)", a.overall_sd],
        ["Individual results", len(a.values)],
        ["Anderson-Darling A²", nm.get("ad")], ["Anderson-Darling p-value", nm.get("p")],
        ["Normal (p ≥ 0.05)", {True: "yes", False: "no"}.get(nm.get("normal"), "n/a")],
        ["Expected out of spec, within σ (ppm)", exp_w.get("total_ppm")],
        ["Expected out of spec, overall s (ppm)", exp_o.get("total_ppm")],
        ["LSL (MPa)", a.lsl], ["USL (MPa)", a.usl], ["Pass rate", p], ["Wilson 95% low", lo_w], ["Wilson 95% high", hi_w],
        ["Out-of-spec builds", len(a.fails)]], numfmt="0.000")

    # --- Signals ---
    sg_ws = wb.create_sheet("Signals")
    nxt = table(sg_ws, 1, ["Variable", "Build ID", "Seq", "Phase", "Value", "Tests", "Reasons"],
                [["Tensile strength", a.records[f.index]["id"], f.index + 1, "baseline" if f.index < a.baseline else "monitoring",
                  a.x[f.index], ", ".join(map(str, f.tests)), "; ".join(f.reasons)] for f in a.flags]
                + [[VARS[v][0], am.records[f.index]["id"], am.idx[f.index] + 1,
                    "baseline" if f.index < am.baseline else "monitoring", am.x[f.index], ", ".join(map(str, f.tests)),
                    "; ".join(f.reasons)] for v, am in ctx.process.items() for f in am.flags],
                [22, 11, 6, 12, 10, 12, 60], numfmt="0.0")
    counts = a.test_counts()
    table(sg_ws, nxt, ["Test (strength)", "Pattern", "Points flagged"],
          [[f"Test {t}", SPECIAL_TESTS[t][1], counts.get(t, 0) if t in a.tests else "off"] for t in SPECIAL_TESTS]
          + [["MR", engine.MR_RULE_TEXT, counts.get("MR", 0) if cfg.get("mr_rule", True) else "off"]])

    # --- Coupons (coupon-level data only) ---
    if xs:
        cp_ws = wb.create_sheet("Coupons")
        cp_rows = []
        for i, r in enumerate(a.records):
            for lab, v in r.get("coupons") or []:
                cp_rows.append([i + 1, r["id"], lab, v,
                                "FAIL" if v < a.lsl or (a.usl is not None and v > a.usl) else "PASS"])
        table(cp_ws, 1, ["Seq", "Build ID", "Coupon", "Strength (MPa)", "Result"], cp_rows, [6, 11, 9, 14, 8], numfmt="0.0")
        for k, row in enumerate(cp_rows, start=2):
            if row[4] == "FAIL":
                cp_ws.cell(row=k, column=5).fill = red_fill
        cp_ws.freeze_panes = "A2"
        cp_ws.auto_filter.ref = f"A1:E{len(cp_rows) + 1}"

    # --- Associations ---
    as_ws = wb.create_sheet("Associations")
    table(as_ws, 1, ["Variable", "Pearson r with strength", "Pairs"],
          [[VARS[x["v"]][0], x["r"], x["n"]] for x in ctx.assoc] or [["(only strength recorded)", None, None]], [24, 22, 8])

    # --- Corrective actions ---
    ca_ws = wb.create_sheet("Corrective actions")
    table(ca_ws, 1, ["Action", "Suggested owner", "Status", "Notes", "Last updated", "Guidance"],
          [[title, owner, ctx.actions.get(tid, {}).get("status", STATUSES_DEFAULT), ctx.actions.get(tid, {}).get("note", ""),
            ctx.actions.get(tid, {}).get("updated", ""), desc] for tid, title, desc, owner in ctx.tasks],
          [32, 18, 20, 40, 16, 80])
    for row in ca_ws.iter_rows(min_row=2, max_col=6):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
